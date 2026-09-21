"""Build a UIT course-project report (.docx) from Markdown, following Nar's report contract.

Vendored from nar-core-space/.claude/skills/word-report/scripts/build_docx.py.
Edit the skill copy first, then copy it back here, so projects do not drift.

Run from the repo root:

    .venv/bin/python scripts/build_report.py
    .venv/bin/python scripts/build_report.py --src report/BAO_CAO.md --out report/BAO_CAO.docx

The contract (see the skill's references/format-contract.md):

- One font family, Times New Roman, everywhere: styles, doc defaults, theme fonts, numbering,
  header, footer, tables, captions and code. Hierarchy comes from size, weight and colour only.
- Page 1 of the file is the table of contents titled by `toc_title`. The cover is printed separately
  from the UIT template, so page numbering starts at `page_number_start` (default 2).
- Header: the course name in Title Case, grey, rule below. Footer: the project title on the left and
  the page number on the right, grey, rule above. Both appear on every page.
- `#` is Heading 1 (20pt), `##` Heading 2 (16pt), `###` Heading 3 (14pt); body and tables 12pt.
- Figures and tables use the full text width; a figure is scaled to fit a 16 x 11.5 cm box so it
  never pushes its caption or the next heading onto another page.

Front matter keys (YAML at the top of the Markdown file):

    course: Xử lý ngôn ngữ tự nhiên          # header, rendered in Title Case
    project_title: <tên đề tài>              # footer, left
    toc_title: Nội dung báo cáo đồ án        # optional, this is the default
    page_number_start: 2                     # optional, this is the default
    header_text: <exact header text>         # optional, overrides the Title Case of `course`
    author: <document author property>       # optional

Other keys (class, lecturer, members, date) are kept for the separate cover and are not printed.

Supported Markdown subset: headings `#` to `####`, paragraphs, bold, italic, inline code, links,
bullet lists, numbered lists (each list restarts at its first number), pipe tables, images,
`>` quotes, fenced code blocks and a line containing only `\\newpage`. HTML comments are dropped.
A missing image is replaced by a red placeholder line and listed at the end; the exit code stays 0.
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

DEFAULT_SRC = "report/BAO_CAO.md"
DEFAULT_OUT = "report/BAO_CAO.docx"

# ------------------------------------------------------------------ contract constants

FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 12
CAPTION_PT = 11
CODE_PT = 11
HEADER_FOOTER_PT = 11
HEADER_FOOTER_COLOR = "7F7F7F"
RULE_COLOR = "7F7F7F"
CODE_SHADE = "F2F2F2"
QUOTE_COLOR = "404040"

# (size pt, colour, space before pt, space after pt)
HEADINGS = {
    1: (20, "365F91", 24, 6),
    2: (16, "4F81BD", 12, 3),
    3: (14, "4F81BD", 10, 2),
    4: (12, "4F81BD", 8, 2),
}

PAGE_W, PAGE_H = Cm(21), Cm(29.7)
MARGIN = Cm(2.5)
TEXT_W = PAGE_W - 2 * MARGIN  # 16 cm
MAX_IMAGE_W = TEXT_W
MAX_IMAGE_H = Cm(11.5)
TWIPS_PER_EMU = 1 / 635

DEFAULT_TOC_TITLE = "Nội dung báo cáo đồ án"
TOC_PLACEHOLDER = "Mục lục chưa cập nhật: mở bằng Word, bấm chuột phải vào đây, chọn Update Field."
CAPTION_PREFIX_TABLE = "Bảng"
CAPTION_MAX_CHARS = 120
LEAD_IN_MAX_CHARS = 200  # a short paragraph right before a table stays on the table's page
KEEP_TABLE_ROWS = 15     # tables up to this many rows are never split across pages
PAGE_BREAK_TOKEN = "\\newpage"

COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\d+)[.)]\s+(.*)$")
QUOTE_RE = re.compile(r"^>\s?(.*)$")
FENCE_RE = re.compile(r"^\s*```")
IMAGE_RE = re.compile(r"^!\[(?P<caption>[^\]]*)\]\((?P<path>[^)]+)\)\s*$")
RULE_RE = re.compile(r"^(\*{3,}|-{3,}|_{3,})$")
TABLE_SEP_RE = re.compile(r"^\|?[\s:|-]*-[\s:|-]*\|?$")
INLINE_RE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<link>\[(?P<ltext>[^\]]*)\]\((?P<lurl>[^)\s]+)\))"
    r"|(?P<bold>\*\*(?P<btext>.+?)\*\*)"
    r"|(?P<italic>\*(?P<itext>[^*]+)\*)"
)
THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
FONT_ATTRS = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")


@dataclass
class InlineRun:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str
    lines: list[str] = field(default_factory=list)
    level: int = 0
    start: int = 1
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    caption: str = ""
    path: str = ""


@dataclass
class BuildStats:
    headings: int = 0
    paragraphs: int = 0
    tables: int = 0
    images: int = 0
    missing_images: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ front matter


def split_front_matter(text: str) -> tuple[dict, str]:
    """Return (front matter dict, body). No front matter gives an empty dict."""
    lines = text.splitlines()
    start = 0
    while start < len(lines) and not lines[start].strip():
        start += 1
    if start >= len(lines) or lines[start].strip() != "---":
        return {}, text
    for idx in range(start + 1, len(lines)):
        if lines[idx].strip() in {"---", "..."}:
            raw = "\n".join(lines[start + 1 : idx])
            return parse_yaml(raw), "\n".join(lines[idx + 1 :])
    return {}, text


def parse_yaml(raw: str) -> dict:
    """Parse the front matter with PyYAML when present, else a flat `key: value` / `- item` subset."""
    try:
        import yaml  # type: ignore

        data = yaml.safe_load(raw) or {}
        return data if isinstance(data, dict) else {}
    except ImportError:
        pass
    data: dict = {}
    current = None
    for line in raw.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        item = re.match(r"^\s+-\s+(.*)$", line)
        if item and current:
            data.setdefault(current, [])
            if isinstance(data[current], list):
                data[current].append(item.group(1).strip())
            continue
        pair = re.match(r"^([A-Za-z_][\w-]*):\s*(.*)$", line)
        if pair:
            current = pair.group(1)
            value = pair.group(2).split(" #")[0].strip()
            data[current] = int(value) if value.isdigit() else (value if value else [])
    return data


def title_case(text: str) -> str:
    """Capitalise the first letter of every word: 'Xử lý ngôn ngữ' -> 'Xử Lý Ngôn Ngữ'."""
    words = []
    for word in text.split():
        words.append(word if word.isupper() else word[:1].upper() + word[1:])
    return " ".join(words)


# ------------------------------------------------------------------ Markdown blocks


def parse_blocks(body: str) -> list[Block]:
    text = COMMENT_RE.sub("", body)
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append(Block(kind="paragraph", lines=list(buffer)))
            buffer.clear()

    idx, total = 0, len(lines)
    while idx < total:
        line = lines[idx].rstrip()
        stripped = line.strip()

        if not stripped:
            flush()
            idx += 1
            continue
        if stripped == PAGE_BREAK_TOKEN:
            flush()
            blocks.append(Block(kind="pagebreak"))
            idx += 1
            continue
        if FENCE_RE.match(line):
            flush()
            idx += 1
            code: list[str] = []
            while idx < total and not FENCE_RE.match(lines[idx]):
                code.append(lines[idx])
                idx += 1
            idx += 1
            blocks.append(Block(kind="code", lines=code))
            continue
        heading = HEADING_RE.match(stripped)
        if heading:
            flush()
            title = heading.group(2).strip().rstrip("#").strip()
            blocks.append(Block(kind="heading", level=len(heading.group(1)), text=title))
            idx += 1
            continue
        if RULE_RE.match(stripped):
            flush()
            idx += 1
            continue
        image = IMAGE_RE.match(stripped)
        if image:
            flush()
            blocks.append(Block(kind="image", caption=image.group("caption").strip(), path=image.group("path").strip()))
            idx += 1
            continue
        if stripped.startswith("|") and idx + 1 < total and TABLE_SEP_RE.match(lines[idx + 1].strip()):
            flush()
            rows = [split_table_row(stripped)]
            idx += 2
            while idx < total and lines[idx].strip().startswith("|"):
                rows.append(split_table_row(lines[idx].strip()))
                idx += 1
            blocks.append(Block(kind="table", rows=rows))
            continue
        quote = QUOTE_RE.match(stripped)
        if quote:
            flush()
            quoted = [quote.group(1).strip()]
            idx += 1
            while idx < total and QUOTE_RE.match(lines[idx].strip()):
                quoted.append(QUOTE_RE.match(lines[idx].strip()).group(1).strip())
                idx += 1
            blocks.append(Block(kind="quote", lines=quoted))
            continue
        bullet = BULLET_RE.match(stripped)
        if bullet:
            flush()
            items = [bullet.group(1).strip()]
            idx += 1
            while idx < total and BULLET_RE.match(lines[idx].strip()):
                items.append(BULLET_RE.match(lines[idx].strip()).group(1).strip())
                idx += 1
            blocks.append(Block(kind="bullets", lines=items))
            continue
        number = NUMBER_RE.match(stripped)
        if number:
            flush()
            items = [number.group(2).strip()]
            first = int(number.group(1))
            idx += 1
            while idx < total and NUMBER_RE.match(lines[idx].strip()):
                items.append(NUMBER_RE.match(lines[idx].strip()).group(2).strip())
                idx += 1
            blocks.append(Block(kind="numbers", lines=items, start=first))
            continue
        buffer.append(stripped)
        idx += 1

    flush()
    return blocks


def split_table_row(line: str) -> list[str]:
    trimmed = line.strip()
    if trimmed.startswith("|"):
        trimmed = trimmed[1:]
    if trimmed.endswith("|"):
        trimmed = trimmed[:-1]
    return [cell.strip() for cell in trimmed.split("|")]


def iter_inline_runs(text: str, bold: bool = False, italic: bool = False) -> Iterator[InlineRun]:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() < pos:
            continue
        if match.start() > pos:
            yield InlineRun(text[pos : match.start()], bold, italic)
        if match.group("code") is not None:
            yield InlineRun(match.group("code")[1:-1], bold, italic, code=True)
        elif match.group("link") is not None:
            label, url = match.group("ltext"), match.group("lurl")
            yield from iter_inline_runs(label, bold, italic)
            if label.strip() != url.strip():
                yield InlineRun(f" ({url})", bold, italic)
        elif match.group("bold") is not None:
            yield from iter_inline_runs(match.group("btext"), True, italic)
        else:
            yield from iter_inline_runs(match.group("itext"), bold, True)
        pos = match.end()
    if pos < len(text):
        yield InlineRun(text[pos:], bold, italic)


def plain_text(text: str) -> str:
    return "".join(piece.text for piece in iter_inline_runs(text))


def add_inline(paragraph, text: str, bold: bool = False, italic: bool = False, size: Pt | None = None) -> None:
    """Pour Markdown inline text into a paragraph. Inline code stays in the report font (one family)."""
    for piece in iter_inline_runs(text, bold, italic):
        if not piece.text:
            continue
        run = paragraph.add_run(piece.text)
        run.bold = piece.bold or None
        run.italic = piece.italic or None
        if size is not None:
            run.font.size = size


# ------------------------------------------------------------------ XML helpers


PPR_ORDER = (
    "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
    "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
    "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange"
).split()
RPR_ORDER = (
    "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
    "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
    "em lang eastAsianLayout specVanish oMath"
).split()


def insert_ordered(parent, child, order: list[str]):
    """Insert `child` where the OOXML schema sequence wants it; Word rejects out-of-order children."""
    local = child.tag.split("}")[1]
    rank = order.index(local)
    for existing in parent:
        tag = existing.tag.split("}")[1] if "}" in existing.tag else existing.tag
        if tag in order and order.index(tag) > rank:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def get_or_insert(parent, tag: str, order: list[str]):
    found = parent.find(qn(f"w:{tag}"))
    return found if found is not None else insert_ordered(parent, OxmlElement(f"w:{tag}"), order)


def set_fonts(rpr, name: str = FONT) -> None:
    """Point every script slot of an rPr at one font and drop theme-font indirection."""
    fonts = get_or_insert(rpr, "rFonts", RPR_ORDER)
    for attr in THEME_FONT_ATTRS:
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in FONT_ATTRS:
        fonts.set(qn(attr), name)


def normalize_fonts(root) -> None:
    """Rewrite every w:rFonts under an XML root to the report font."""
    for fonts in root.iter(qn("w:rFonts")):
        for attr in THEME_FONT_ATTRS:
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]
        for attr in FONT_ATTRS:
            fonts.set(qn(attr), FONT)


def border(side: str, color: str = RULE_COLOR, size: int = 4, space: int = 1):
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    return el


def paragraph_border(paragraph, side: str) -> None:
    pbdr = get_or_insert(paragraph._p.get_or_add_pPr(), "pBdr", PPR_ORDER)
    pbdr.append(border(side))


def shade(paragraph, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(paragraph._p.get_or_add_pPr(), shd, PPR_ORDER)


def set_tabs(ppr, stops: list[tuple[str, int]], leader: str | None = None) -> None:
    """Replace a pPr's tab stops. stops = [(alignment, position in twips)]."""
    old = ppr.find(qn("w:tabs"))
    if old is not None:
        ppr.remove(old)
    tabs = OxmlElement("w:tabs")
    for align, pos in stops:
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), align)
        if leader and align == "right":
            tab.set(qn("w:leader"), leader)
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)
    insert_ordered(ppr, tabs, PPR_ORDER)


def add_field(paragraph, instr: str, result: str = "", dirty: bool = False) -> None:
    """Append a complex field (begin / instr / separate / result / end) to a paragraph."""
    def fld(kind: str):
        run = OxmlElement("w:r")
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        if kind == "begin" and dirty:
            char.set(qn("w:dirty"), "true")
        run.append(char)
        paragraph._p.append(run)

    fld("begin")
    run = OxmlElement("w:r")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instr} "
    run.append(text)
    paragraph._p.append(run)
    fld("separate")
    if result:
        paragraph.add_run(result)
    fld("end")


def text_width_twips() -> int:
    return int(TEXT_W * TWIPS_PER_EMU)


# ------------------------------------------------------------------ styles and page setup


def style_or_none(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def ensure_paragraph_style(doc, style_id: str, name: str, based_on: str = "Normal"):
    """Create a paragraph style by raw XML so Word recognises built-in names such as 'toc 1'."""
    styles = doc.styles.element
    for existing in styles.findall(qn("w:style")):
        if existing.get(qn("w:styleId")) == style_id:
            return existing
    style = OxmlElement("w:style")
    style.set(qn("w:type"), "paragraph")
    style.set(qn("w:styleId"), style_id)
    for tag, val in (("w:name", name), ("w:basedOn", based_on), ("w:next", "Normal"), ("w:uiPriority", "39")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        style.append(el)
    style.append(OxmlElement("w:unhideWhenUsed"))
    styles.append(style)
    return style


def style_ppr(style_el):
    ppr = style_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        rpr = style_el.find(qn("w:rPr"))
        (rpr.addprevious(ppr) if rpr is not None else style_el.append(ppr))
    return ppr


def style_rpr(style_el):
    rpr = style_el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        style_el.append(rpr)
    return rpr


def setup_styles(doc) -> None:
    styles_el = doc.styles.element

    # Document defaults: one font, 12pt, Vietnamese proofing, 1.15 line spacing.
    defaults = styles_el.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(rpr_default)
    for tag in ("sz", "szCs"):
        get_or_insert(rpr_default, tag, RPR_ORDER).set(qn("w:val"), str(BODY_PT * 2))
    get_or_insert(rpr_default, "lang", RPR_ORDER).set(qn("w:val"), "vi-VN")

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    set_fonts(normal.element.get_or_add_rPr())
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for level, (size, color, before, after) in HEADINGS.items():
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        set_fonts(style.element.get_or_add_rPr())

    toc_heading = style_or_none(doc, "TOC Heading")
    if toc_heading is not None:
        toc_heading.paragraph_format.space_before = Pt(0)
        toc_heading.paragraph_format.space_after = Pt(12)

    # TOC entry styles; Word fills them when it updates the field.
    for level, indent, before, bold in ((1, 0, 6, True), (2, 240, 0, False), (3, 480, 0, False)):
        style = ensure_paragraph_style(doc, f"TOC{level}", f"toc {level}")
        ppr = style_ppr(style)
        for child in list(ppr):
            ppr.remove(child)
        set_tabs(ppr, [("right", text_width_twips())], leader="dot")
        spacing = get_or_insert(ppr, "spacing", PPR_ORDER)
        spacing.set(qn("w:before"), str(before * 20))
        spacing.set(qn("w:after"), "40")
        if indent:
            get_or_insert(ppr, "ind", PPR_ORDER).set(qn("w:left"), str(indent))
        rpr = style_rpr(style)
        for child in list(rpr):
            rpr.remove(child)
        if bold:
            insert_ordered(rpr, OxmlElement("w:b"), RPR_ORDER)

    for name in ("Header", "Footer"):
        style = style_or_none(doc, name)
        if style is None:
            continue
        style.font.size = Pt(HEADER_FOOTER_PT)
        style.font.color.rgb = RGBColor.from_string(HEADER_FOOTER_COLOR)
        set_tabs(style_ppr(style.element), [("center", text_width_twips() // 2), ("right", text_width_twips())])

    caption = style_or_none(doc, "Caption")
    if caption is not None:
        caption.font.size = Pt(CAPTION_PT)
        caption.font.bold = False
        caption.font.italic = True
        caption.font.color.rgb = RGBColor(0, 0, 0)

    # Drop theme-font indirection from every style, then point every slot at the report font.
    normalize_fonts(styles_el)


def setup_numbering_fonts(doc) -> None:
    """Bullets and numbers use the report font too: Symbol/Wingdings bullets become plain glyphs."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:  # noqa: BLE001 - template without numbering
        return
    glyphs = {"": "•", "o": "–", "": "•", "": "•"}
    for lvl in numbering.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        text = lvl.find(qn("w:lvlText"))
        if fmt is not None and fmt.get(qn("w:val")) == "bullet" and text is not None:
            text.set(qn("w:val"), glyphs.get(text.get(qn("w:val")), "•"))
        # A level without rPr lets Word pick a fallback font for the number and its tab.
        if lvl.find(qn("w:rPr")) is None:
            lvl.append(OxmlElement("w:rPr"))
        set_fonts(lvl.find(qn("w:rPr")))
    normalize_fonts(numbering)


def setup_theme(doc) -> None:
    """Set the theme's major and minor Latin fonts to the report font, so no fallback can differ."""
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            xml = part.blob.decode("utf-8")
            xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', rf"\g<1>{FONT}\g<2>", xml)
            part._blob = xml.encode("utf-8")


def setup_settings(doc) -> None:
    """Ask Word to refresh fields (the TOC) when the file is opened."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def setup_section(doc, meta: dict) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.top_margin = section.bottom_margin = MARGIN
    section.left_margin = section.right_margin = MARGIN
    section.header_distance = section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False

    sect_pr = section._sectPr
    start = int(meta.get("page_number_start", 2))
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        cols = sect_pr.find(qn("w:cols"))
        cols.addprevious(pg_num) if cols is not None else sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))

    header_text = meta.get("header_text") or title_case(str(meta.get("course", "")).strip())
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.style = doc.styles["Header"]
    hp.add_run(header_text)
    paragraph_border(hp, "bottom")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.style = doc.styles["Footer"]
    ppr = fp._p.get_or_add_pPr()
    set_tabs(ppr, [("clear", text_width_twips() // 2), ("right", text_width_twips())])
    paragraph_border(fp, "top")
    fp.add_run(str(meta.get("project_title", "")).strip())
    fp.add_run("\t")
    add_field(fp, "PAGE", str(start))

    normalize_fonts(header._element)
    normalize_fonts(footer._element)


# ------------------------------------------------------------------ content blocks


def add_toc_page(doc, meta: dict) -> None:
    title = doc.add_paragraph(str(meta.get("toc_title") or DEFAULT_TOC_TITLE), style="TOC Heading")
    title.paragraph_format.keep_with_next = True
    field_par = doc.add_paragraph(style=next(s for s in doc.styles if s.style_id == "TOC1"))
    add_field(field_par, 'TOC \\o "1-3" \\h \\z \\u', TOC_PLACEHOLDER, dirty=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def fit_image(width: int, height: int) -> tuple[int, int]:
    """Scale (up or down) to the largest size inside the MAX_IMAGE_W x MAX_IMAGE_H box."""
    ratio = min(MAX_IMAGE_W / width, MAX_IMAGE_H / height)
    return int(width * ratio), int(height * ratio)


def add_picture_block(doc, block: Block, base_dir: Path, stats: BuildStats, figure_no: int) -> int:
    target = (base_dir / block.path).resolve()
    if not target.exists():
        stats.missing_images.append(block.path)
        print(f"  Cảnh báo: thiếu hình {block.path}")
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"[Thiếu hình: {block.path}]")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return figure_no

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    picture = paragraph.add_run().add_picture(str(target))
    picture.width, picture.height = (Emu(v) for v in fit_image(picture.width, picture.height))
    stats.images += 1

    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    label = f"Hình {figure_no}. {plain_text(block.caption)}" if block.caption else f"Hình {figure_no}."
    caption.add_run(label)
    return figure_no + 1


def column_widths(rows: list[list[str]], n_cols: int) -> list[int]:
    """Split the text width across columns by content length, with a floor so short columns stay legible."""
    total = text_width_twips()
    weights = []
    for c in range(n_cols):
        longest = max((len(plain_text(row[c])) if c < len(row) else 0) for row in rows)
        weights.append(min(max(longest, 6), 60))
    floor = int(Cm(1.6) * TWIPS_PER_EMU)
    raw = [total * w / sum(weights) for w in weights]
    widths = [max(int(w), floor) for w in raw]
    excess = sum(widths) - total
    if excess > 0:
        flexible = [i for i, w in enumerate(widths) if w > floor]
        pool = sum(widths[i] - floor for i in flexible) or 1
        for i in flexible:
            widths[i] -= int(excess * (widths[i] - floor) / pool)
    widths[-1] += total - sum(widths)
    return widths


def add_table_block(doc, block: Block, stats: BuildStats) -> None:
    rows = block.rows
    if not rows:
        return
    n_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    if style_or_none(doc, "Table Grid") is not None:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    widths = column_widths(rows, n_cols)
    for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(width))

    keep_together = len(rows) <= KEEP_TABLE_ROWS
    for r_idx, row in enumerate(rows):
        tr_pr = table.rows[r_idx]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))  # a row never breaks across pages
        if r_idx == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))  # the header row repeats on the next page
        for c_idx in range(n_cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Emu(int(widths[c_idx] / TWIPS_PER_EMU))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(paragraph, value, bold=(r_idx == 0), size=Pt(TABLE_PT))
            # Short tables stay on one page: every row but the last keeps with the next.
            if keep_together and r_idx < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
    stats.tables += 1


def add_quote_block(doc, block: Block, stats: BuildStats) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(1)
    paragraph.paragraph_format.right_indent = Cm(1)
    add_inline(paragraph, " ".join(line for line in block.lines if line), italic=True)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(QUOTE_COLOR)
    stats.paragraphs += 1


def add_code_block(doc, block: Block) -> None:
    """Code stays in the report font; the grey band and indent mark it as code."""
    for line in block.lines or [""]:
        paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_after = fmt.space_before = Pt(0)
        fmt.line_spacing = 1.0
        fmt.left_indent = Cm(0.5)
        run = paragraph.add_run(line if line else " ")
        run.font.size = Pt(CODE_PT)
        shade(paragraph, CODE_SHADE)


def restart_numbering(doc, paragraphs: list, start: int) -> None:
    """Give one numbered list its own w:num so it restarts at `start` instead of continuing."""
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"].element
    num_pr = style.find(qn("w:pPr")).find(qn("w:numPr"))
    base_num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == base_num_id:
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    if abstract_id is None:
        return
    new_id = str(max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), new_id)
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_el = OxmlElement("w:startOverride")
    start_el.set(qn("w:val"), str(start))
    override.append(start_el)
    num.append(override)
    numbering.append(num)
    for paragraph in paragraphs:
        ppr = paragraph._p.get_or_add_pPr()
        pnum = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), new_id)
        pnum.append(ilvl)
        pnum.append(num_id)
        old = ppr.find(qn("w:numPr"))
        if old is not None:
            ppr.remove(old)
        insert_ordered(ppr, pnum, PPR_ORDER)


def is_table_caption(block: Block, next_block: Block | None) -> bool:
    if next_block is None or next_block.kind != "table":
        return False
    text = " ".join(block.lines).strip()
    return text.startswith(CAPTION_PREFIX_TABLE) and len(text) < CAPTION_MAX_CHARS


def build_document(meta: dict, blocks: list[Block], base_dir: Path, stats: BuildStats):
    doc = Document()
    setup_styles(doc)
    setup_numbering_fonts(doc)
    setup_theme(doc)
    setup_settings(doc)
    setup_section(doc, meta)
    add_toc_page(doc, meta)

    figure_no = 1
    previous = ""
    for position, block in enumerate(blocks):
        next_block = blocks[position + 1] if position + 1 < len(blocks) else None

        if block.kind == "heading":
            stats.headings += 1
            paragraph = doc.add_heading(level=min(block.level, 4))
            add_inline(paragraph, block.text)
        elif block.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif block.kind == "paragraph":
            text = " ".join(block.lines).strip()
            if text:
                stats.paragraphs += 1
                paragraph = doc.add_paragraph()
                if previous == "table":
                    paragraph.paragraph_format.space_before = Pt(8)
                if is_table_caption(block, next_block):
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.keep_with_next = True
                    add_inline(paragraph, text, bold=True)
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    if next_block is not None and next_block.kind == "table" and len(text) <= LEAD_IN_MAX_CHARS:
                        paragraph.paragraph_format.keep_with_next = True
                    add_inline(paragraph, text)
        elif block.kind == "bullets":
            for item in block.lines:
                add_inline(doc.add_paragraph(style="List Bullet"), item)
                stats.paragraphs += 1
        elif block.kind == "numbers":
            items = []
            for item in block.lines:
                paragraph = doc.add_paragraph(style="List Number")
                add_inline(paragraph, item)
                items.append(paragraph)
                stats.paragraphs += 1
            restart_numbering(doc, items, block.start)
        elif block.kind == "table":
            add_table_block(doc, block, stats)
        elif block.kind == "image":
            figure_no = add_picture_block(doc, block, base_dir, stats, figure_no)
        elif block.kind == "quote":
            add_quote_block(doc, block, stats)
        elif block.kind == "code":
            add_code_block(doc, block)
        previous = block.kind

    # Nothing added in the body may reintroduce another font.
    normalize_fonts(doc.element.body)

    props = doc.core_properties
    props.title = str(meta.get("project_title", ""))
    props.subject = title_case(str(meta.get("course", "")))
    props.author = str(meta.get("author", ""))
    props.last_modified_by = str(meta.get("author", ""))
    props.comments = ""
    props.created = props.modified = datetime.now()
    return doc


def build_report(src: Path, out: Path) -> BuildStats:
    if not src.exists():
        raise SystemExit(f"Không tìm thấy file nguồn: {src}")
    meta, body = split_front_matter(src.read_text(encoding="utf-8"))
    for key in ("course", "project_title"):
        if not meta.get(key):
            raise SystemExit(f"Front matter thiếu khóa bắt buộc `{key}` trong {src}")
    stats = BuildStats()
    doc = build_document(meta, parse_blocks(body), src.parent, stats)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return stats


def resolve(path_arg: str) -> Path:
    path = Path(path_arg).expanduser()
    return path if path.is_absolute() else (ROOT / path)


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Dựng báo cáo Word từ Markdown theo chuẩn báo cáo đồ án UIT.")
    parser.add_argument("--src", default=DEFAULT_SRC, help=f"File Markdown nguồn (mặc định {DEFAULT_SRC}).")
    parser.add_argument("--out", default=DEFAULT_OUT, help=f"File .docx kết quả (mặc định {DEFAULT_OUT}).")
    args = parser.parse_args(list(argv) if argv is not None else None)

    src, out = resolve(args.src), resolve(args.out)
    print(f"Đang dựng báo cáo từ {src}")
    stats = build_report(src, out)
    print("\nTóm tắt lần dựng:")
    print(f"  Tiêu đề      : {stats.headings}")
    print(f"  Đoạn văn     : {stats.paragraphs}")
    print(f"  Bảng         : {stats.tables}")
    print(f"  Hình         : {stats.images}")
    print(f"  Hình bị thiếu: {len(stats.missing_images)}")
    for path in stats.missing_images:
        print(f"      - {path}")
    print(f"  Kết quả      : {out}  ({out.stat().st_size / 1024:.1f} KB)")
    print("  Mục lục cần Word cập nhật (Update Field) trước khi nộp.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
